import os
import io
import zipfile
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from PIL import Image

def escapar_caracteres_latex(texto):
    especiais = {
        '&': r'\&',
        '%': r'\%',
        '$': r'\$',
        '#': r'\#',
        '_': r'\_',
        '{': r'\{',
        '}': r'\}',
        '~': r'\textasciitilde{}',
        '^': r'\textasciicircum{}',
        '\\': r'\textbackslash{}',
    }
    for char, escape in especiais.items():
        texto = texto.replace(char, escape)
    return texto

def converter_pptx_para_latex(caminho_arquivo):
    if not os.path.exists(caminho_arquivo):
        print(f"O arquivo {caminho_arquivo} não foi encontrado.")
        return

    prs = Presentation(caminho_arquivo)
    latex_output = []

    downloads_path = os.path.join(os.path.expanduser("~"), "Downloads")
    output_folder = os.path.join(downloads_path, "output")
    os.makedirs(output_folder, exist_ok=True)

    nome_arquivo = os.path.splitext(os.path.basename(caminho_arquivo))[0]
    output_tex = os.path.join(output_folder, f"{nome_arquivo}.tex")

    latex_output.extend([
        r"\documentclass{beamer}",
        r"\usepackage[utf8]{inputenc}",
        r"\usepackage{graphicx}",
        r"\usetheme{Madrid}",
        "",
        r"\begin{document}",
        ""
    ])

    for slide_idx, slide in enumerate(prs.slides):
        latex_output.append(r"\begin{frame}")

        titulo = None
        itens_slide = []
        imagens_embutidas = []

        for shape_idx, shape in enumerate(slide.shapes):
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                image = shape.image
                image_bytes = image.blob
                image_stream = io.BytesIO(image_bytes)

                pil_image = Image.open(image_stream).convert("RGB")

                imagem_nome = f"{nome_arquivo}_slide{slide_idx}_img{shape_idx}.png"
                imagem_path = os.path.join(output_folder, imagem_nome)
                pil_image.save(imagem_path)

                imagens_embutidas.append(imagem_nome)

            elif hasattr(shape, "text_frame") and shape.has_text_frame:
                parags = shape.text_frame.paragraphs

                if not titulo:
                    if parags and parags[0].text.strip():
                        titulo = parags[0].text.strip()

                    for parag in parags[1:]:
                        if parag.text.strip():
                            itens_slide.append(parag)
                else:
                    for parag in parags:
                        if parag.text.strip():
                            itens_slide.append(parag)

        if titulo:
            latex_output.append(f"\\frametitle{{{escapar_caracteres_latex(titulo)}}}")

        if itens_slide:
            latex_output.append(r"\begin{itemize}")
            for parag in itens_slide:
                texto = escapar_caracteres_latex(parag.text.strip())
                indent = '  ' * parag.level  
                latex_output.append(f"{indent}\\item {texto}")
            latex_output.append(r"\end{itemize}")

        for imagem_nome in imagens_embutidas:
            latex_output.append(r"\begin{center}")
            latex_output.append(fr"\includegraphics[width=0.8\textwidth]{{{imagem_nome}}}")
            latex_output.append(r"\end{center}")

        latex_output.append(r"\end{frame}")
        latex_output.append("")

    latex_output.append(r"\end{document}")

    with open(output_tex, "w", encoding="utf-8") as f:
        f.write("\n".join(latex_output))

    print(f"Arquivo LaTeX salvo em: {output_tex}")

    zip_path = output_tex.replace(".tex", ".zip")
    with zipfile.ZipFile(zip_path, "w") as zipf:
        zipf.write(output_tex, arcname=os.path.basename(output_tex))

        for imagem_nome in os.listdir(output_folder):
            if imagem_nome.endswith(".png") and imagem_nome.startswith(nome_arquivo):
                imagem_path = os.path.join(output_folder, imagem_nome)
                zipf.write(imagem_path, arcname=imagem_nome)

    print(f"Arquivo ZIP gerado em: {zip_path}")
    return zip_path
